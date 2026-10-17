"""Reporte Topográfico (ST.CA.PC.001.F001).

Lo hace el cadista, lo valida un supervisor de Calidad (si está mal, vuelve al cadista) y queda disponible para
unirse, uno a uno, a un protocolo. Viaja dentro del PDF del protocolo (reporte + plano) y KP lo firma junto con él.

Casi todo lo llena el sistema: actividades fijas, georreferencia del proyecto, equipos (del Excel "Log de Rep de
Insp de los EME_TOPOGRAFIA") y los datos de campo (del CSV del levantamiento: punto, norte, este, elevación).
"""
import csv
import io
import json
import re
from datetime import date, datetime

import openpyxl

from . import almacen, config, correo, db
from . import logica as L
from .logica import CALIDAD, TOPOGRAFICO, ErrorNegocio, _evento, _exigir, _inst, _tipo

GEO_DEFECTO = {"punto": "BASE H.B.P.", "tipo": "UTM", "este": "200966.792", "norte": "8398587.436", "elevacion": "4211.390"}
MAX_PUNTOS = 16
EXTENSIONES_LEVANTAMIENTO = (".csv", ".txt", ".xlsx")


def es_topo(t):
    return t["codigo"] == TOPOGRAFICO


def _sin_extension(nombre):
    return re.sub(r"\.[A-Za-z0-9]{2,4}$", "", nombre or "")


def _fecha(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v).strip()


# ---------------------------------------------------------------------------
# Configuración (equipos, topógrafos, georreferencia): la mantiene un administrador
# ---------------------------------------------------------------------------
def georreferencia():
    guardada = db._meta("georreferencia")
    return {**GEO_DEFECTO, **(json.loads(guardada) if guardada else {})}


def equipos():
    """Equipos vigentes del Excel. El nivel automático no se usa en los reportes."""
    filas = db.filas("select * from equipo where activo = 1 order by orden, modelo")
    hoy = date.today().isoformat()
    return [{**e, "vencido": bool(e["proxima_calibracion"]) and e["proxima_calibracion"] < hoy}
            for e in filas if "NIVEL" not in (e["descripcion"] or "").upper()]


def topografos():
    return [x["nombre"] for x in db.filas("select nombre from topografo order by orden, nombre")]


def configuracion():
    return {"equipos": equipos(), "topografos": topografos(), "georreferencia": georreferencia()}


def importar_equipos(yo, datos):
    """Reemplaza los equipos con el Excel del año (los que ya no están quedan inactivos; los reportes hechos no cambian)."""
    L._exigir_admin(yo)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(datos), data_only=True, read_only=True)
    except Exception:
        raise ErrorNegocio("Suba el Excel de equipos (.xlsx): Log de Rep de Insp de los EME_TOPOGRAFIA")
    lista, cab = [], None
    for fila in wb.worksheets[0].iter_rows(values_only=True):
        textos = [str(c or "").strip().upper() for c in fila]
        if cab is None:
            if "MODELO" in textos and any("SERIE" in x for x in textos):
                cab = textos
            continue
        d = dict(zip(cab, fila))
        col = lambda pal: next((v for k, v in d.items() if pal in k), None)  # noqa: E731
        modelo, serie = _fecha(d.get("MODELO")), _fecha(col("SERIE"))
        if not modelo or not serie:
            continue
        lista.append({"serie": serie, "rieme": _fecha(col("RIEME")), "descripcion": _fecha(col("EQUIPOS")),
                      "marca": _fecha(d.get("MARCA")), "modelo": modelo, "certificado": _fecha(col("CERTIFICADO")),
                      "fecha_calibracion": _fecha(d.get("FECHA DE CALIBRACIÓN")),
                      "proxima_calibracion": _fecha(col("PRÓXIMA CALIBRACIÓN"))})
    if not lista:
        raise ErrorNegocio("No se encontraron equipos: revise que sea el Log de Rep de Insp de los EME (columnas MODELO y N° SERIE)")
    with db.transaccion():
        db.ejecutar("update equipo set activo = 0")
        for n, e in enumerate(lista):
            db.ejecutar("""insert into equipo (serie, rieme, descripcion, marca, modelo, certificado, fecha_calibracion,
                           proxima_calibracion, activo, orden) values (?, ?, ?, ?, ?, ?, ?, ?, 1, ?)
                           on conflict(serie) do update set rieme = excluded.rieme, descripcion = excluded.descripcion,
                           marca = excluded.marca, modelo = excluded.modelo, certificado = excluded.certificado,
                           fecha_calibracion = excluded.fecha_calibracion, proxima_calibracion = excluded.proxima_calibracion,
                           activo = 1, orden = excluded.orden""",
                         e["serie"], e["rieme"], e["descripcion"], e["marca"], e["modelo"], e["certificado"],
                         e["fecha_calibracion"], e["proxima_calibracion"], n)
    return {"equipos": len(equipos())}


def agregar_topografo(yo, nombre):
    L._exigir_admin(yo)
    nombre = re.sub(r"\s+", " ", nombre or "").strip()
    if not nombre:
        raise ErrorNegocio("Escriba el nombre del topógrafo")
    n = (db.fila("select count(*) as n from topografo") or {}).get("n", 0)
    db.ejecutar("insert into topografo (nombre, orden) values (?, ?) on conflict(nombre) do nothing", nombre, n)
    return topografos()


def quitar_topografo(yo, nombre):
    L._exigir_admin(yo)
    db.ejecutar("delete from topografo where nombre = ?", nombre)
    return topografos()


def guardar_georreferencia(yo, datos):
    """Fija para el proyecto. Solo cambia si el sistema se usa en otra unidad."""
    L._exigir_admin(yo)
    geo = {"punto": str(datos.get("punto") or "").strip(), "tipo": "UTM"}
    if not geo["punto"]:
        raise ErrorNegocio("Indique el punto B.M.")
    for k in ("este", "norte", "elevacion"):
        try:
            geo[k] = f"{float(str(datos.get(k)).replace(',', '.')):.3f}"
        except (TypeError, ValueError):
            raise ErrorNegocio(f"{k.capitalize()}: debe ser un número")
    db._meta("georreferencia", json.dumps(geo))
    return geo


# ---------------------------------------------------------------------------
# Lo que el sistema llena solo (al guardar y al enviar): nadie lo cambia desde la pantalla
# ---------------------------------------------------------------------------
def _equipo_fila(serie):
    e = db.fila("select * from equipo where serie = ?", serie) if serie else None
    return e and {"marca": e["marca"], "modelo": e["modelo"], "serie": e["serie"], "calibracion": "SI",
                  "certificado": e["certificado"], "_descripcion": e["descripcion"] or ""}


def _plano_de(i):
    a = db.fila("select nombre_archivo from adjunto where instancia_id = ? and version = ? and ref_schema = 'plano' "
                "order by created_at desc", i["id"], i["version_actual"])
    return _sin_extension(a["nombre_archivo"]) if a else ""


def normalizar(i, t, respuestas):
    r = {"checklist": {}, "campos": {}, "tablas": {}, "adjuntos": {}, "descripcion": "", "observaciones": "",
         **(respuestas or {})}
    if not es_topo(t):
        return aplicar_vinculo(i, t, r)
    previo = i.get("respuestas") or {}
    for s in t["schema"]["secciones"]:
        if s["tipo"] == "checklist" and s.get("fijo"):
            r["checklist"] = {it["id"]: {"r": it.get("valor", ""), "obs": "", "extra": ""}
                              for g in s["grupos"] for it in g["items"]}
    campos, tablas = {**r["campos"]}, {**r["tablas"]}
    filas_previas = (previo.get("tablas") or {}).get("equipos") or []
    filas = []
    for k, n in (("equipo_1", 0), ("equipo_2", 1)):
        fila = _equipo_fila(campos.get(k))
        if campos.get(k) and not fila:              # equipo dado de baja después: se conserva lo que ya tenía
            fila = filas_previas[n] if n < len(filas_previas) else {}
        filas.append({k2: v for k2, v in (fila or {}).items() if not k2.startswith("_")})
        if n == 0:
            desc = (fila or {}).get("_descripcion", "")
            campos["equipo_medicion"] = (f"{'GPS ' if 'GNSS' in desc.upper() else ''}{fila['marca']} {fila['modelo']} "
                                         f"{fila['serie']}").strip() if fila else ""
    tablas["equipos"] = filas
    tablas["georreferencia"] = [georreferencia()]
    # Datos de campo y nombre del archivo: solo los cambia la carga del CSV
    tablas["puntos_referencia"] = (previo.get("tablas") or {}).get("puntos_referencia") or []
    campos["archivo"] = (previo.get("campos") or {}).get("archivo") or ""
    campos["plano_adjunto"] = "SI"
    r.update(campos=campos, tablas=tablas, observaciones="")
    r["adjuntos"] = {**r["adjuntos"], "plano": {"adjunta": "SI", "numero": _plano_de(i)}}
    return r


def aplicar_vinculo(i, t, r):
    """En un protocolo: el reporte topográfico unido (y su plano, que va en el mismo PDF) se marca solo."""
    defs = t["schema"].get("adjuntos") or []
    fuente = next((a for a in defs if a.get("fuente") == "reporte_topo"), None)
    if not fuente:
        return r
    adj, campos = {**(r.get("adjuntos") or {})}, {**(r.get("campos") or {})}
    rep = db.fila("select * from instancia where vinculado_a = ?", i["id"])
    rt_campo = any(c.get("fuente") == "reporte_topo" for s in t["schema"]["secciones"] for c in s.get("campos") or [])
    if rep:
        adj[fuente["id"]] = {"adjunta": "SI", "numero": rep["correlativo"], "reporte_id": rep["id"]}
        plano = _plano_de(rep)
        for a in defs:
            if a.get("incluido_en_reporte"):
                adj[a["id"]] = {"adjunta": "SI", "numero": plano or f"incluido en el Reporte Topográfico {rep['correlativo']}",
                                "en_reporte": True}
        if rt_campo:
            campos["rt_numero"] = rep["correlativo"]
    else:
        if (adj.get(fuente["id"]) or {}).get("reporte_id"):
            adj[fuente["id"]] = {"adjunta": "", "numero": ""}
        for a in defs:
            if (adj.get(a["id"]) or {}).get("en_reporte"):
                adj[a["id"]] = {"adjunta": "", "numero": ""}
        if rt_campo and (r.get("campos") or {}).get("rt_numero"):
            campos["rt_numero"] = ""
    return {**r, "adjuntos": adj, "campos": campos}


# ---------------------------------------------------------------------------
# Datos de campo: CSV del levantamiento (punto, norte, este, elevación; lo que sigue se descarta)
# ---------------------------------------------------------------------------
def _numero(v):
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).strip().replace(",", ".")) if str(v or "").strip() else None


def leer_levantamiento(nombre, datos):
    if nombre.lower().endswith(".xlsx"):
        try:
            ws = openpyxl.load_workbook(io.BytesIO(datos), data_only=True, read_only=True).worksheets[0]
        except Exception:
            raise ErrorNegocio("No se pudo leer el archivo del levantamiento")
        crudas = [list(f) for f in ws.iter_rows(values_only=True)]
    else:
        try:
            texto = datos.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = datos.decode("cp1252", errors="replace")
        try:
            dialecto = csv.Sniffer().sniff(texto[:2000], delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        crudas = list(csv.reader(io.StringIO(texto), dialecto))
    puntos = []
    for n, f in enumerate(crudas, start=1):
        if not any(str(c or "").strip() for c in f):
            continue
        if len(f) < 4:
            raise ErrorNegocio(f"Fila {n}: se esperan 4 columnas (punto, norte, este, elevación)")
        try:
            norte, este, elev = _numero(f[1]), _numero(f[2]), _numero(f[3])
        except ValueError:
            if not puntos:                      # primera fila con títulos
                continue
            raise ErrorNegocio(f"Fila {n}: norte, este y elevación deben ser números")
        if None in (norte, este, elev):
            raise ErrorNegocio(f"Fila {n}: faltan coordenadas")
        if norte < este:
            raise ErrorNegocio(f"Fila {n}: el norte debe ser el número más alto (orden: punto, norte, este, elevación)")
        punto = f[0]
        punto = str(int(punto)) if isinstance(punto, float) and punto.is_integer() else str(punto or "").strip()
        puntos.append({"punto": punto, "tipo": "UTM", "este": f"{este:.3f}", "norte": f"{norte:.3f}", "elevacion": f"{elev:.3f}"})
    if not puntos:
        raise ErrorNegocio("El archivo no tiene puntos")
    if len(puntos) > MAX_PUNTOS:
        raise ErrorNegocio(f"El archivo tiene {len(puntos)} puntos: el reporte admite como máximo {MAX_PUNTOS}")
    return puntos


def subir_levantamiento(yo, id_, nombre, datos):
    if not nombre.lower().endswith(EXTENSIONES_LEVANTAMIENTO):
        raise ErrorNegocio("Suba el CSV del levantamiento (.csv, .txt o .xlsx)")
    puntos = leer_levantamiento(nombre, datos)
    with db.transaccion():
        i = _inst(id_)
        t = _tipo(i["tipo_id"])
        _exigir(es_topo(t), "Solo el Reporte Topográfico lleva datos de campo")
        L._editable(i)
        L._exigir_elaborar(yo, i["tipo_id"])
        for a in db.filas("select * from adjunto where instancia_id = ? and version = ? and ref_schema = 'levantamiento'",
                          id_, i["version_actual"]):
            db.ejecutar("delete from adjunto where id = ?", a["id"])
            almacen.borrar(a["ruta"])
        seguro = re.sub(r"[^\w.-]+", "_", nombre)
        ruta = L._guardar_archivo(f"adjuntos/{id_}/V{i['version_actual']}", f"{db.nuevo_id()}-{seguro}", datos)
        db.insertar("adjunto", {"id": db.nuevo_id(), "instancia_id": id_, "version": i["version_actual"], "tipo": "levantamiento",
                                "ref_schema": "levantamiento", "nombre_archivo": nombre, "ruta": ruta, "tamano": len(datos),
                                "subido_por": yo["id"], "created_at": db.ahora()})
        r = {**(i["respuestas"] or {})}
        r["tablas"] = {**(r.get("tablas") or {}), "puntos_referencia": puntos}
        r["campos"] = {**(r.get("campos") or {}), "archivo": _sin_extension(nombre)}
        i["respuestas"] = r
        db.actualizar("instancia", id_, {"respuestas": normalizar(i, t, r), "updated_at": db.ahora()})
    return {"puntos": len(puntos), "archivo": _sin_extension(nombre)}


# ---------------------------------------------------------------------------
# Index del cadista, validación del supervisor y unión a un protocolo
# ---------------------------------------------------------------------------
def _resumen(i):
    t = _tipo(i["tipo_id"])
    firmas = {f["rol"]: f for f in db.filas("select * from firma where instancia_id = ? and version = ?",
                                            i["id"], i["version_actual"])}
    prot = db.fila("select id, correlativo from instancia where id = ?", i["vinculado_a"]) if i.get("vinculado_a") else None
    creador = db.fila("select nombre from perfil where id = ?", i["creado_por"]) if i["creado_por"] else None
    rev = firmas.get("revisado_aprobado") or {}
    obs = db.fila("select detalle from evento where instancia_id = ? and tipo = 'observado_topo' order by created_at desc",
                  i["id"])
    campos = (i["respuestas"] or {}).get("campos") or {}
    return {
        "id": i["id"], "correlativo": i["correlativo"], "estado": i["estado"],
        "numero_registro": L._numero_registro(t, i["correlativo"]) if i["correlativo"] else None,
        "fecha": i["cabecera"].get("fecha"), "entregable": i["cabecera"].get("entregable"),
        "estructura": i["cabecera"].get("estructura"), "descripcion": (i["respuestas"] or {}).get("descripcion") or "",
        "topografo": campos.get("topografo") or "", "equipo": campos.get("equipo_medicion") or "",
        "plano": _plano_de(i), "cadista": (firmas.get("elaborado") or {}).get("nombre") or (creador or {}).get("nombre"),
        "validado_por": rev.get("nombre") if rev.get("estado") == "firmado" else None,
        "protocolo": prot, "observacion": (obs["detalle"] or {}).get("motivo") if obs and i["estado"] == "observado" else None,
        "updated_at": i["updated_at"],
    }


def reportes():
    filas = db.filas("select * from instancia where tipo_id = ? order by created_at desc", TOPOGRAFICO)
    return [_resumen(L._inst(f["id"])) for f in filas]


def disponibles(id_):
    """Reportes validados que aún no están unidos a ningún protocolo (y el que ya está unido a este)."""
    filas = db.filas("select id from instancia where tipo_id = ? and ((estado = 'disponible' and vinculado_a is null) "
                     "or vinculado_a = ?) order by correlativo", TOPOGRAFICO, id_)
    return [_resumen(_inst(f["id"])) for f in filas]


def revisar(yo, id_, decision, motivo, contexto=None):
    """El supervisor de Calidad que corresponda valida (firma "Revisado") u observa (vuelve al cadista)."""
    _exigir(yo["rol"] in CALIDAD, "Solo un supervisor de Calidad valida los reportes topográficos")
    if decision not in ("aceptado", "observado"):
        raise ErrorNegocio("Decisión inválida")
    motivo = (motivo or "").strip()
    if decision == "observado" and not motivo:
        raise ErrorNegocio("Indique qué debe corregir el cadista")
    with db.transaccion():
        i = _inst(id_)
        if not es_topo(_tipo(i["tipo_id"])) or i["estado"] != "revision_topo":
            raise ErrorNegocio("Este reporte no está esperando validación")
        cadista = db.fila("select * from perfil where id = ?", i["creado_por"]) if i["creado_por"] else None
        if decision == "aceptado":
            f = db.fila("select * from firma where instancia_id = ? and version = ? and rol = 'revisado_aprobado'",
                        id_, i["version_actual"])
            db.actualizar("firma", f["id"], {"perfil_id": yo["id"], "nombre": yo["nombre"], "cargo": yo.get("cargo"),
                                             "email": yo["email"], "estado": "firmado", "firmado_at": db.ahora(),
                                             "evidencia": {"metodo": "Sesión en la app (personal Stracon)", **(contexto or {})}})
            db.actualizar("instancia", id_, {"estado": "disponible", "updated_at": db.ahora()})
            _evento(id_, "validado_topo", {}, yo["email"])
            _evento(id_, "firmado", {"rol": "revisado_aprobado"}, yo["email"])
            aviso = f"{yo['nombre']} validó el Reporte Topográfico {i['correlativo']}: ya puede unirse a un protocolo."
        else:
            db.actualizar("instancia", id_, {"estado": "observado", "updated_at": db.ahora()})
            _evento(id_, "observado_topo", {"motivo": motivo}, yo["email"])
            aviso = f"{yo['nombre']} observó el Reporte Topográfico {i['correlativo']}: {motivo}. Corríjalo y vuelva a enviarlo."
        if cadista:
            correo.enviar([cadista["email"]], f"[Stracon] Reporte Topográfico {i['correlativo']}: "
                          f"{'validado' if decision == 'aceptado' else 'observado'}", [aviso],
                          {"texto": "Abrir reporte", "url": f"{config.URL_APP}/p/{id_}"})
    if decision == "aceptado":
        L._intentar_pdf(id_)                 # reporte + plano en un solo PDF
    return {"estado": "disponible" if decision == "aceptado" else "observado"}


def vincular(yo, id_, reporte_id):
    """Une (o suelta, con reporte_id vacío) el reporte topográfico de un protocolo. Uno a uno."""
    with db.transaccion():
        i = _inst(id_)
        t = _tipo(i["tipo_id"])
        L._editable(i)
        L._exigir_elaborar(yo, i["tipo_id"])
        if not any(a.get("fuente") == "reporte_topo" for a in t["schema"].get("adjuntos") or []):
            raise ErrorNegocio("Este formato no lleva reporte topográfico")
        actual = db.fila("select id, correlativo from instancia where vinculado_a = ?", id_)
        if actual and actual["id"] != reporte_id:
            db.actualizar("instancia", actual["id"], {"vinculado_a": None, "estado": "disponible", "updated_at": db.ahora()})
            _evento(actual["id"], "desvinculado", {"protocolo": i["correlativo"] or "en borrador"}, yo["email"])
        if reporte_id and (not actual or actual["id"] != reporte_id):
            rep = _inst(reporte_id)
            if not es_topo(_tipo(rep["tipo_id"])):
                raise ErrorNegocio("Elija un Reporte Topográfico")
            if rep["vinculado_a"]:
                raise ErrorNegocio(f"El Reporte Topográfico {rep['correlativo']} ya está unido a otro protocolo")
            if rep["estado"] != "disponible":
                raise ErrorNegocio(f"El Reporte Topográfico {rep['correlativo']} aún no fue validado por un supervisor")
            db.actualizar("instancia", reporte_id, {"vinculado_a": id_, "estado": "vinculado", "updated_at": db.ahora()})
            _evento(reporte_id, "vinculado", {"protocolo": i["correlativo"] or "en borrador"}, yo["email"])
        db.actualizar("instancia", id_, {"respuestas": aplicar_vinculo(i, t, i["respuestas"]), "updated_at": db.ahora()})
    return {"ok": True}


def reporte_de(protocolo_id):
    return db.fila("select * from instancia where vinculado_a = ?", protocolo_id)
