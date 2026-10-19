"""QC Index (Control de Registros de Calidad, ST.CA.PR.005.F001) generado desde los protocolos.

Mismas hojas y columnas que el Excel que hoy llena Calidad a mano. Las columnas propias de cada formato
se definen en protocolos/tipos.json → "qc_index": {"hoja", "columnas": [[título, fuente], ...]}.
Fuentes: "c:<campo>" campo del formato · "k:<ítem>" dato extra de un ítem del checklist · "turno" ·
"prog:<campo>" progresiva 0+028.87 → 28.87 · "calc:longitud" progresiva final − inicial.
"""
import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

from . import config, db
from .logica import _numero_registro, maestros, proyecto, tipos

MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Set", "Oct", "Nov", "Dic"]
INICIO = ["ÍTEM", "FECHA", "CÓDIGO DEL FORMATO", "NOMBRE DEL FORMATO", "N° DE REGISTRO",
          "DESCRIPCIÓN DEL TRABAJO / ACTIVIDAD", "ENTREGABLE", "ESTRUCTURA", "ELEMENTO"]
AZUL = PatternFill("solid", fgColor="0B5CAD")
CELESTE = PatternFill("solid", fgColor="DCE8F5")
FINO = Side(style="thin", color="8A96A3")
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)


def _progresiva(v):
    """'0+028.87' → 28.87 · '2+039.90' → 2039.9 · número → número."""
    if v in (None, ""):
        return None
    s = str(v).strip().replace(",", "")
    m = re.fullmatch(r"(\d+)\+(\d+(?:\.\d+)?)", s)
    if m:
        return round(int(m.group(1)) * 1000 + float(m.group(2)), 3)
    try:
        return float(s)
    except ValueError:
        return s


def _valor(fuente, valor):
    r = valor.get("respuestas") or {}
    campos = r.get("campos") or {}
    if fuente == "turno":
        return f"Turno {campos['turno']}" if campos.get("turno") else "-"
    if fuente == "calc:longitud":
        a, b = _progresiva(campos.get("prog_inicial")), _progresiva(campos.get("prog_final"))
        return round(b - a, 3) if isinstance(a, float) and isinstance(b, float) else "-"
    tipo, _, clave = fuente.partition(":")
    if tipo == "c":
        v = campos.get(clave)
    elif tipo == "k":
        v = ((r.get("checklist") or {}).get(clave) or {}).get("extra")
    elif tipo == "prog":
        v = _progresiva(campos.get(clave))
    else:
        v = None
    return "-" if v in (None, "") else v


def _estado(inst, conforme):
    if inst["estado"] == "anulado":
        return "Anulado"
    if conforme is False:
        return "No Conforme"
    return "Conforme" if inst["estado"] == "aprobado" else "Revisión"


def _adjuntos(tipo, valor):
    marcados = (valor.get("respuestas") or {}).get("adjuntos") or {}
    lineas = [f".{a['label']} {marcados[a['id']].get('numero') or ''}".strip()
              for a in tipo["schema"].get("adjuntos") or [] if (marcados.get(a["id"]) or {}).get("adjunta") == "SI"]
    return "\n".join(lineas) or "-"


# Metadatos de los formatos del QC Index (iguales a las plantillas de Calidad)
COD_RESUMEN, COD_DETALLE, VERSION_FMT, FECHA_FMT = "ST.CA.PR.005.F001", "ST.CA.PR.005.F002", "Versión: 02", "Fecha: 22/09/2023"
MESES_LARGOS = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SETIEMBRE", "OCTUBRE",
                "NOVIEMBRE", "DICIEMBRE"]
NEGRITA = Font(bold=True)


def _fin_de_mes(d):
    return date(d.year + d.month // 12, d.month % 12 + 1, 1).toordinal() - 1


def etiqueta_periodo(desde=None, hasta=None):
    """'OCTUBRE 2026' (un mes completo) · '01/10/2026 al 15/10/2026' · 'AL 19/10/2026' (todo el contrato)."""
    def f(d):
        return f"{d.day:02d}/{d.month:02d}/{d.year}"
    if desde and hasta and desde.day == 1 and (desde.year, desde.month) == (hasta.year, hasta.month) \
            and hasta.toordinal() == _fin_de_mes(hasta):
        return f"{MESES_LARGOS[desde.month - 1]} {desde.year}"
    if desde and hasta:
        return f"{f(desde)} al {f(hasta)}"
    if desde:
        return f"DESDE {f(desde)}"
    return f"AL {f(hasta or date.today())}"


def _encabezado(ws, titulo, codigo, ncols, proy, periodo, disciplina=None):
    """Como la plantilla: logo, título, código/versión/fecha del formato, proyecto, cliente, disciplina y mes/año."""
    ncols = max(ncols, 8)
    t0, t1, c0, c1 = 3, ncols - 2, ncols - 1, ncols
    ws.merge_cells(start_row=1, start_column=1, end_row=3, end_column=2)
    ruta_logo = config.WEB / "img" / "logo.png"
    if ruta_logo.exists():
        logo = XLImage(str(ruta_logo))
        logo.width, logo.height = int(42 * logo.width / logo.height), 42
        ws.add_image(logo, "A1")
    ws.merge_cells(start_row=1, start_column=t0, end_row=2, end_column=t1)
    ws.cell(1, t0, titulo).font = Font(bold=True, size=14)
    ws.cell(1, t0).alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=3, start_column=t0, end_row=3, end_column=t1)
    ws.cell(3, t0, "FORMATO").alignment = Alignment(horizontal="center")
    for fila, texto in ((1, codigo), (2, VERSION_FMT), (3, FECHA_FMT)):
        ws.merge_cells(start_row=fila, start_column=c0, end_row=fila, end_column=c1)
        ws.cell(fila, c0, texto).font = NEGRITA if fila == 1 else Font()
    ws.cell(4, 1, "NOMBRE DEL PROYECTO: ").font = NEGRITA
    ws.cell(4, 3, proy["nombre_contrato"])
    ws.cell(4, t1, "N° REGISTRO:").font = NEGRITA
    ws.merge_cells(start_row=4, start_column=c0, end_row=4, end_column=c1)
    ws.cell(4, c0, "-")
    ws.cell(5, 1, "CLIENTE:").font = NEGRITA
    ws.cell(5, 3, proy["cliente"])
    if disciplina is not None:
        ws.cell(5, 6, "DISCIPLINA:").font = NEGRITA
        ws.cell(5, 7, disciplina or "-")
    ws.cell(5, t1, "MES / AÑO:").font = NEGRITA
    ws.merge_cells(start_row=5, start_column=c0, end_row=5, end_column=c1)
    ws.cell(5, c0, periodo)
    for fila in range(1, 6):
        for col in range(1, ncols + 1):
            ws.cell(fila, col).border = BORDE


def _cabecera_tabla(ws, titulos):
    ws.append(titulos)
    fila = ws.max_row
    for celda in ws[fila]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = AZUL
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = BORDE
    ws.row_dimensions[fila].height = 32
    ws.freeze_panes = f"A{fila + 1}"
    return fila


def _anchos(ws, titulos):
    anchos = {"ÍTEM": 6, "FECHA": 11, "CÓDIGO DEL FORMATO": 26, "NOMBRE DEL FORMATO": 26, "N° DE REGISTRO": 38,
              "DESCRIPCIÓN DEL TRABAJO / ACTIVIDAD": 60, "DOCUMENTOS ADJUNTOS": 34,
              "DOCUMENTOS DE REFERENCIA (PLANOS REF.)": 30, "OBSERVACIONES / COMENTARIOS": 30}
    for i, t in enumerate(titulos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = anchos.get(t, 16)


def filas_tipo(t, entregable=None, tipos_sel=None, desde=None, hasta=None, solo_aprobados=False):
    """Protocolos de un formato que entran con los filtros (la fecha es la del protocolo), ordenados por fecha."""
    if tipos_sel and t["codigo"] not in tipos_sel:
        return []
    filas = db.filas("""select i.*, v.snapshot from instancia i
                        join version v on v.instancia_id = i.id and v.numero = i.version_actual
                        where i.tipo_id = ? and i.correlativo is not null order by i.created_at""", t["id"])
    salida = []
    for f in filas:
        cab = (f["snapshot"] or {}).get("cabecera") or {}
        if entregable and cab.get("entregable") != entregable:
            continue
        if solo_aprobados and f["estado"] != "aprobado":
            continue
        try:
            fecha = date.fromisoformat(cab.get("fecha") or "")
        except ValueError:
            fecha = None
        if (desde and (not fecha or fecha < desde)) or (hasta and (not fecha or fecha > hasta)):
            continue
        salida.append((f, fecha))
    salida.sort(key=lambda x: (x[1] or date.max, x[0]["correlativo"] or ""))
    return salida


def generar(entregable=None, tipos_sel=None, desde=None, hasta=None, solo_aprobados=False):
    """Devuelve el Excel (bytes) con los datos de la versión vigente de cada protocolo con correlativo.
    Filtros: entregable ("RC"), formatos (códigos), rango de fechas del protocolo y solo aprobados (dossier)."""
    proy = proyecto()
    ents = {e["codigo"]: e["nombre"] for e in maestros()["entregables"]}
    periodo = etiqueta_periodo(desde, hasta)
    wb = Workbook()
    resumen = wb.active
    resumen.title = COD_RESUMEN
    totales = []

    for t in tipos():
        q = t.get("qc_index")
        if not q or (tipos_sel and t["codigo"] not in tipos_sel):
            continue
        titulos = INICIO + [c[0] for c in q["columnas"]]
        if not q.get("sin_adjuntos"):
            titulos.append("DOCUMENTOS ADJUNTOS")
        titulos.append("DOCUMENTOS DE REFERENCIA (PLANOS REF.)")
        if not q.get("sin_originador"):
            titulos.append("ORIGINADOR")
        titulos += ["ESTADO", "OBSERVACIONES / COMENTARIOS"]

        ws = wb.create_sheet(q["hoja"])
        _encabezado(ws, "QC INDEX", COD_DETALLE, len(titulos), proy, periodo, q.get("disciplina", "-"))
        _cabecera_tabla(ws, titulos)
        _anchos(ws, titulos)

        conteo = {"total": 0, "conformes": 0, "no_conformes": 0, "meses": [0] * 12}
        for n, (i, fecha) in enumerate(filas_tipo(t, entregable, tipos_sel, desde, hasta, solo_aprobados), start=1):
            valor = i["snapshot"]
            cab = valor.get("cabecera") or {}
            conforme = valor.get("conforme")
            elaborado = db.fila("select nombre from firma where instancia_id = ? and version = ? and rol = 'elaborado'",
                                i["id"], i["version_actual"])
            planos = "\n".join([*(cab.get("plano_ref") or []), *([cab["plano_otro"]] if cab.get("plano_otro") else [])]) or "-"
            fila = [n, fecha, f"{t['codigo']} [{proy['codigo_contrato']}]", t["nombre"], _numero_registro(t, i["correlativo"]),
                    (valor.get("respuestas") or {}).get("descripcion") or "-", ents.get(cab.get("entregable"), cab.get("entregable")),
                    cab.get("estructura") or "-", cab.get("elemento") or "-"]
            fila += [_valor(c[1], valor) for c in q["columnas"]]
            if not q.get("sin_adjuntos"):
                fila.append(_adjuntos(t, valor))
            fila.append(planos)
            if not q.get("sin_originador"):
                fila.append(elaborado["nombre"] if elaborado else "-")
            estado = _estado(i, conforme)
            fila += [estado, (valor.get("respuestas") or {}).get("observaciones") or "-"]
            ws.append(fila)
            for celda in ws[ws.max_row]:
                celda.border = BORDE
                celda.alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row=ws.max_row, column=2).number_format = "DD/MM/YYYY"

            conteo["total"] += 1
            conteo["conformes"] += estado == "Conforme"
            conteo["no_conformes"] += estado == "No Conforme"
            if fecha:
                conteo["meses"][fecha.month - 1] += 1
        totales.append((t, conteo))

    # Hoja resumen (ST.CA.PR.005.F001): totales, conformes, no conformes, ICRC y registros por mes
    titulos = ["CÓDIGO DEL FORMATO", "NOMBRE DEL FORMATO", "TOTAL DE REGISTROS", "CONFORMES", "NO CONFORMES", "ICRC", "", *MESES]
    _encabezado(resumen, "CONTROL DE REGISTROS DE CALIDAD", COD_RESUMEN, len(titulos), proy, periodo)
    if entregable:
        resumen.cell(5, 6, "ENTREGABLE:").font = NEGRITA
        resumen.cell(5, 7, ents.get(entregable, entregable))
    _cabecera_tabla(resumen, titulos)
    for t, c in totales:
        icrc = round(c["conformes"] / c["total"], 2) if c["total"] else None
        resumen.append([f"{t['codigo']} [{proy['codigo_contrato']}]", t["nombre"], c["total"], c["conformes"],
                        c["no_conformes"], icrc, "", *c["meses"]])
        for celda in resumen[resumen.max_row]:
            celda.border = BORDE
    total = [sum(c[k] for _, c in totales) for k in ("total", "conformes", "no_conformes")]
    resumen.append(["", "TOTAL", *total, round(total[1] / total[0], 2) if total[0] else None, "",
                    *[sum(c["meses"][m] for _, c in totales) for m in range(12)]])
    for celda in resumen[resumen.max_row]:
        celda.font = Font(bold=True)
        celda.fill = CELESTE
        celda.border = BORDE
    resumen.column_dimensions["A"].width = 30
    resumen.column_dimensions["B"].width = 42
    for col in range(3, len(titulos) + 1):
        resumen.column_dimensions[get_column_letter(col)].width = 12 if col <= 6 else 6

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
