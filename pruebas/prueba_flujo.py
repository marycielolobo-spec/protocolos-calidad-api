"""Prueba automática del flujo completo, sin tocar los datos reales (usa una carpeta temporal).

Uso:  .venv\\Scripts\\python pruebas\\prueba_flujo.py
"""
import base64
import io
import os
import re
import sys
import tempfile
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")
os.environ["PROTOCOLOS_DATOS"] = tempfile.mkdtemp(prefix="protocolos-prueba-")
# Siempre local: base SQLite temporal, archivos locales, correos simulados (aunque .env tenga Supabase o Gmail)
for _k in ("DATABASE_URL", "SUPABASE_URL", "SUPABASE_PUBLISHABLE_KEY", "SUPABASE_SECRET_KEY",
           "SMTP_USER", "SMTP_PASSWORD", "RESEND_API_KEY"):
    os.environ[_k] = ""
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402
from pypdf import PdfReader  # noqa: E402

from servidor.main import app  # noqa: E402

c = TestClient(app)
pasos = 0


def ok(cond, msg):
    global pasos
    if not cond:
        raise AssertionError(msg)
    pasos += 1
    print(f"  OK  {msg}")


def api(metodo, ruta, usuario=None, esperado=200, **kw):
    r = c.request(metodo, f"/api{ruta}", headers={"X-Usuario": usuario} if usuario else {}, **kw)
    if r.status_code != esperado:
        raise AssertionError(f"{metodo} {ruta}: esperado {esperado}, llegó {r.status_code} {r.text[:300]}")
    return r.json() if "json" in r.headers.get("content-type", "") else r


def texto_pdf(datos):
    return "".join(p.extract_text() for p in PdfReader(io.BytesIO(datos)).pages)


def llenar(tipo):
    """Respuestas válidas para cualquier formato."""
    s = tipo["schema"]
    checklist, campos = {}, {}
    for sec in s["secciones"]:
        if sec["tipo"] == "checklist":
            positivo = "C" if sec["opciones"] == "C_NC_NA" else "SI"
            for g in sec["grupos"]:
                for it in g["items"]:
                    checklist[it["id"]] = {"extra": "x"} if it.get("solo_extra") else {"r": positivo}
        if sec["tipo"] == "campos":
            for cp in sec["campos"]:
                campos[cp["id"]] = cp["opciones"][0] if cp.get("tipo") == "radio" else 1 if cp.get("tipo") == "number" else "x"
    campos.update({"prog_inicial": "0+028.87", "prog_final": "0+048.38", "area": 440.306})
    adjuntos = {a["id"]: {"adjunta": "NA"} for a in s.get("adjuntos") or []}
    return {"checklist": checklist, "campos": campos, "tablas": {}, "adjuntos": adjuntos,
            "descripcion": "Se liberó terreno de fundación (prueba)", "observaciones": ""}


def cabecera(d, entregable, estructura, planos=()):
    return {**d["inst"]["cabecera"], "entregable": entregable, "estructura": estructura, "plano_ref": list(planos)}


def firma_png():
    img = Image.new("RGBA", (300, 100), (255, 255, 255, 0))
    ImageDraw.Draw(img).line([(10, 70), (80, 20), (150, 80), (290, 30)], fill=(0, 0, 0, 255), width=4)
    b = io.BytesIO()
    img.save(b, "PNG")
    return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()


def ultimo_codigo():
    asunto = next(x["asunto"] for x in api("GET", "/correos") if "Código para firmar" in x["asunto"])
    return re.search(r"(\d{6})", asunto).group(1)


print("Flujo completo de protocolos")
ses = api("GET", "/sesion")
por_rol = {p["rol"]: p["id"] for p in ses["perfiles"]}
sup, jefe, ops, ana, cad = (por_rol[r] for r in ("supervisor_calidad", "jefe_calidad", "operaciones", "control_documentario",
                                                  "topografia"))
c.headers["X-Usuario"] = sup        # consultas sin usuario explícito: como el supervisor de Calidad
kp = ses["contactos"][0]["id"]
tipos = api("GET", "/tipos")
ok(len(tipos) == 9 and all(t["sigla"] for t in tipos), "9 formatos cargados (con Geoceldas), cada uno con su sigla")
m = api("GET", "/maestros")
rc = next(e for e in m["entregables"] if e["codigo"] == "RC")
tmf = next(e for e in m["entregables"] if e["codigo"] == "TMF")
ok(len(m["entregables"]) == 4 and len(tmf["estructuras"]) == 6 and rc["planos"], "Maestros: 4 entregables (con WRF), 6 estructuras en TMF, planos de RC")
ok("Relleno Común" in m["listas"]["material_relleno"], "Listas para desplegables cargadas")

fund = next(t for t in tipos if t["codigo"] == "ST.CA.PC.001.F001".replace("F001", "F002"))   # Terreno de fundación
otro = next(t for t in tipos if t["codigo"] == "ST.CA.PC.008.F002")                          # Drenes
A = api("POST", "/protocolos", sup, json={"tipo_id": fund["id"], "origen": "solicitud_operaciones"})["id"]
B = api("POST", "/protocolos", sup, json={"tipo_id": otro["id"], "origen": "backlog"})["id"]
C = api("POST", "/protocolos", sup, json={"tipo_id": fund["id"], "origen": "solicitud_operaciones"})["id"]
firmantes = {"firmantes": {"elaborado": {"perfil_id": sup}, "revisado_aprobado": {"perfil_id": ops}}}

r = api("POST", f"/protocolos/{A}/enviar-firma", sup, 400, json=firmantes)
ok("Entregable" in r["error"], "No deja enviar sin entregable ni datos")

# Estructura que no corresponde al entregable
d = api("GET", f"/protocolos/{A}")
resp_a = llenar(fund)
adj = next(a_ for a_ in fund["schema"]["adjuntos"] if a_["id"] == "plano_geologico")
resp_a["adjuntos"][adj["id"]] = {"adjunta": "SI", "numero": "RT-001"}
api("PUT", f"/protocolos/{A}/valor", sup, json={"cabecera": cabecera(d, "RC", tmf["estructuras"][0]["nombre"]),
                                              "respuestas": resp_a, "conforme": True})
r = api("POST", f"/protocolos/{A}/enviar-firma", sup, 400, json=firmantes)
ok("no pertenece" in r["error"], "Estructura de otro entregable: lo detiene")

cab_a = cabecera(d, "RC", "Reservorio Cunahuiri", rc["planos"][:2])
api("PUT", f"/protocolos/{A}/valor", sup, json={"cabecera": cab_a, "respuestas": resp_a, "conforme": True})
r = api("POST", f"/protocolos/{A}/enviar-firma", sup, 400, json=firmantes)
ok("falta subir" in r["error"], 'Marcó "Sí se adjunta" sin archivo: lo detiene')

png = io.BytesIO()
Image.new("RGB", (400, 300), "steelblue").save(png, "PNG")
api("POST", f"/protocolos/{A}/adjuntos", sup, files={"archivo": ("plano.png", png.getvalue(), "image/png")},
    data={"tipo": adj["tipo_adjunto"], "ref_schema": adj["id"]})

resp_neg = llenar(fund)
resp_neg["adjuntos"] = resp_a["adjuntos"]
primero = next(k for k, v in resp_neg["checklist"].items() if "r" in v)
resp_neg["checklist"][primero]["r"] = "NO"
api("PUT", f"/protocolos/{A}/valor", sup, json={"cabecera": cab_a, "respuestas": resp_neg, "conforme": True})
r = api("POST", f"/protocolos/{A}/enviar-firma", sup, 400, json=firmantes)
ok("requiere observación" in r["error"], "Respuesta NO sin observación: lo detiene")

api("PUT", f"/protocolos/{A}/valor", sup, json={"cabecera": cab_a, "respuestas": resp_a, "conforme": True})

# Reporte Topográfico: el cadista lo hace → un supervisor lo valida → se une (uno a uno) al protocolo A
topo_t = next(t_ for t_ in tipos if t_["codigo"] == "ST.CA.PC.001.F001")
conf = api("GET", "/topografia/config", cad)
ok(conf["georreferencia"]["punto"] == "BASE H.B.P." and conf["topografos"] and conf["equipos"]
   and not any("NIVEL" in (e["descripcion"] or "") for e in conf["equipos"]),
   "Topografía: georreferencia fija, equipos del Excel (sin nivel automático) y topógrafos")
R = api("POST", "/protocolos", cad, json={"tipo_id": topo_t["id"], "origen": "solicitud_operaciones"})["id"]
d = api("GET", f"/protocolos/{R}", cad)
ok(d["inst"]["respuestas"]["checklist"]["c3"]["r"] == "NO" and d["inst"]["respuestas"]["checklist"]["c8"]["r"] == "NA",
   "Actividades fijas: las pone el sistema")
gnss = next(e for e in conf["equipos"] if e["modelo"] == "R980")
api("PUT", f"/protocolos/{R}/valor", cad, json={
    "cabecera": cabecera(d, "RC", "Reservorio Cunahuiri", rc["planos"][:1]), "conforme": None,
    "respuestas": {**d["inst"]["respuestas"], "descripcion": "Levantamiento de fundación (prueba)", "observaciones": "no va",
                   "checklist": {"c1": {"r": "NO"}},
                   "campos": {"equipo_1": gnss["serie"], "topografo": conf["topografos"][0]}}})
rr = api("GET", f"/protocolos/{R}", cad)["inst"]["respuestas"]
ok(rr["checklist"]["c1"]["r"] == "SI" and rr["observaciones"] == "" and rr["tablas"]["georreferencia"][0]["este"] == "200966.792",
   "Nadie cambia actividades, observaciones ni georreferencia")
eq = rr["tablas"]["equipos"][0]
ok(eq["marca"] == "TRIMBLE" and eq["serie"] == gnss["serie"] and eq["calibracion"] == "SI" and eq["certificado"] == gnss["certificado"]
   and rr["campos"]["equipo_medicion"] == f"GPS TRIMBLE R980 {gnss['serie']}",
   "Se elige el equipo y marca, serie, certificado y calibración salen solos")
csv_largo = "\n".join(f"{n},8398516.903,203316.548,4227.918" for n in range(1, 18)).encode()
r = api("POST", f"/protocolos/{R}/levantamiento", cad, 400, files={"archivo": ("largo.csv", csv_largo, "text/csv")})
ok("16" in r["error"], "CSV con más de 16 puntos: rechazado")
csv_ok = "1,8398516.903,203316.548,4227.918,\n2,8398522.877,203315.201,4228.077,504.6\n".encode()
api("POST", f"/protocolos/{R}/levantamiento", cad, files={"archivo": ("260925-C2-R COMUN.csv", csv_ok, "text/csv")})
rr = api("GET", f"/protocolos/{R}", cad)["inst"]["respuestas"]
pr = rr["tablas"]["puntos_referencia"]
ok(len(pr) == 2 and pr[0]["norte"] == "8398516.903" and pr[0]["este"] == "203316.548" and pr[1]["elevacion"] == "4228.077"
   and pr[0]["tipo"] == "UTM" and rr["campos"]["archivo"] == "260925-C2-R COMUN",
   "CSV: punto, norte, este y elevación (lo demás se descarta); el archivo se llama como el CSV")
r = api("POST", f"/protocolos/{R}/enviar-firma", cad, 400, json={"firmantes": {}})
ok("falta subir" in r["error"], "Sin plano no se envía: el plano es obligatorio")
api("POST", f"/protocolos/{R}/adjuntos", cad, files={"archivo": ("2611-2026-PTF-RC-TS-001.png", png.getvalue(), "image/png")},
    data={"tipo": "plano", "ref_schema": "plano"})
api("POST", f"/protocolos/{R}/enviar-firma", cad, json={"firmantes": {}})
d = api("GET", f"/protocolos/{R}", cad)
ok(d["inst"]["estado"] == "revision_topo" and d["inst"]["correlativo"] == "001" and d["numero_registro"].endswith("] - 001"),
   "Cadista firma y envía: numeración propia (001) y pasa al supervisor")
ok(api("GET", f"/protocolos/{A}/reportes-disponibles", sup) == [], "Sin validar no aparece para unirse a un protocolo")
api("POST", f"/protocolos/{R}/revision-topo", cad, 403, json={"decision": "aceptado"})
api("POST", f"/protocolos/{R}/revision-topo", sup, json={"decision": "observado", "motivo": "Falta el punto 3"})
ok(api("GET", f"/protocolos/{R}", cad)["inst"]["estado"] == "observado", "El supervisor observa: vuelve al cadista")
api("POST", f"/protocolos/{R}/enviar-firma", cad, json={"firmantes": {}})
api("POST", f"/protocolos/{R}/revision-topo", sup, json={"decision": "aceptado"})
d = api("GET", f"/protocolos/{R}", cad)
ok(d["inst"]["estado"] == "disponible" and d["firmas"][1]["estado"] == "firmado" and d["firmas"][1]["perfil_id"] == sup,
   "El supervisor que lo toma lo valida: queda su firma y el reporte disponible")
r = c.get(f"/api/protocolos/{R}/pdf")
ok(r.status_code == 200 and len(PdfReader(io.BytesIO(r.content)).pages) >= 2 and "BASE H.B.P." in texto_pdf(r.content),
   "Reporte y plano en un solo PDF")
ok([x["id"] for x in api("GET", f"/protocolos/{A}/reportes-disponibles", sup)] == [R], "El protocolo ve el reporte validado y libre")
api("POST", f"/protocolos/{A}/reporte-topo", sup, json={"reporte_id": R})
ra = api("GET", f"/protocolos/{A}")["inst"]["respuestas"]
ok(ra["adjuntos"]["reporte_topografico"]["numero"] == "001" and ra["adjuntos"]["plano_utm"]["numero"] == "2611-2026-PTF-RC-TS-001"
   and ra["campos"]["rt_numero"] == "001", "Al unirlo, el reporte N° 001 y su plano quedan marcados solos")
api("PUT", f"/protocolos/{A}/valor", sup, json={"cabecera": cab_a, "respuestas": resp_a, "conforme": True})
ok(api("GET", f"/protocolos/{A}")["inst"]["respuestas"]["adjuntos"]["reporte_topografico"]["reporte_id"] == R,
   "Guardar el protocolo no suelta el reporte unido")
ok(api("GET", f"/protocolos/{R}", cad)["inst"]["estado"] == "vinculado" and api("GET", f"/protocolos/{C}/reportes-disponibles", sup) == [],
   "Unido a A: ya no aparece para otro protocolo")
r = api("POST", f"/protocolos/{C}/reporte-topo", sup, 400, json={"reporte_id": R})
ok("otro protocolo" in r["error"], "Un reporte, un protocolo")
api("GET", "/topografia/reportes", cad)
ok(any(x["id"] == R and x["protocolo"] for x in api("GET", "/topografia/reportes", cad)), "Index topográfico: muestra a qué protocolo se unió")
api("POST", f"/protocolos/{A}/enviar-firma", sup, json=firmantes)
for X, t_, ent, est in ((B, otro, "TMF", tmf["estructuras"][0]["nombre"]), (C, fund, "RC", "Reservorio Cunahuiri")):
    d = api("GET", f"/protocolos/{X}")
    api("PUT", f"/protocolos/{X}/valor", sup, json={"cabecera": cabecera(d, ent, est), "respuestas": llenar(t_), "conforme": True})
    api("POST", f"/protocolos/{X}/enviar-firma", sup, json=firmantes)
a = api("GET", f"/protocolos/{A}")
ok(a["inst"]["correlativo"] == "ATF-RC-001", f"Correlativo con sigla y entregable: {a['inst']['correlativo']}")
ok(a["numero_registro"] == "ST.CA.PC.001.F002 [10203002] - ATF-RC-001", f"N° de registro como el QC Index: {a['numero_registro']}")
ok(api("GET", f"/protocolos/{C}")["inst"]["correlativo"] == "ATF-RC-002", "El número sigue por formato y entregable (ATF-RC-002)")
ok(api("GET", f"/protocolos/{B}")["inst"]["correlativo"] == "ID-001", "En TMF el correlativo no lleva código de entregable (ID-001)")
ok(a["firmas"][0]["fecha_documento"] == cab_a["fecha"], "La firma del elaborador lleva la fecha del protocolo")
api("PUT", f"/protocolos/{A}/valor", sup, 400, json={"cabecera": cab_a, "respuestas": resp_a, "conforme": True})
ok(True, "En firma ya no se puede editar")

# Firmar y enviar: el supervisor firma al enviar → Control Documentario valida → Operaciones firma
ok(a["inst"]["estado"] == "validacion_doc" and a["firmas"][0]["estado"] == "firmado" and a["firmas"][0]["perfil_id"] == sup,
   "Firmar y enviar: queda la firma del supervisor y pasa a Control Documentario")
ok(any("Validar ATF-RC-001" in x["asunto"] for x in api("GET", "/correos")), "Al analista le llega el aviso")
f_rev = a["firmas"][1]["id"]
r = api("POST", f"/firmas/{f_rev}/firmar", ops, 400)
ok("firma interna" in r["error"], "Operaciones no firma antes de la validación documentaria")
val = api("GET", "/validacion", ana)
fila_a = next(x for x in val["protocolos"] if x["id"] == A)
ok(len(val["protocolos"]) == 3 and all(x["ok"] for x in fila_a["automatica"]),
   "Bandeja del analista con la revisión automática en verde")
api("GET", "/validacion", ops, 403)
api("POST", f"/protocolos/{A}/validacion", sup, 403, json={"decision": "aceptado", "checks": []})
ok(True, "Solo el analista valida")
todos = [x["id"] for x in val["checks"]]
r = api("POST", f"/protocolos/{A}/validacion", ana, 400, json={"decision": "aceptado", "checks": todos[:1]})
ok("Marque todos" in r["error"], "Para validar debe marcar toda la lista de verificación")
api("POST", f"/protocolos/{A}/validacion", ana, 400, json={"decision": "observado", "motivo": ""})
api("POST", f"/protocolos/{A}/validacion", ana, json={"decision": "observado", "motivo": "Falta el número del RT"})
ok(api("GET", f"/protocolos/{A}")["inst"]["estado"] == "observado", "Observado por el analista → vuelve al supervisor")
api("POST", f"/protocolos/{A}/enviar-firma", sup, json=firmantes)
fila_a = next(x for x in api("GET", "/validacion", ana)["protocolos"] if x["id"] == A)
ok(fila_a["observaciones_previas"][0]["motivo"] == "Falta el número del RT",
   "Corregido y reenviado: vuelve al analista con su observación previa")
for X in (A, B, C):
    api("POST", f"/protocolos/{X}/validacion", ana, json={"decision": "aceptado", "checks": todos})
a = api("GET", f"/protocolos/{A}")
ok(a["inst"]["estado"] == "firma_interna", "Validado → pasa a Operaciones")
f_rev = a["firmas"][1]["id"]
api("POST", f"/firmas/{f_rev}/firmar", sup, 403)
ok(True, "Nadie firma por otra persona")
ok(api("POST", f"/firmas/{f_rev}/firmar", ops)["estado"] == "listo_revision", "A listo para revisión tras la firma de Operaciones")
for X in (B, C):
    x = api("GET", f"/protocolos/{X}")
    api("POST", f"/firmas/{x['firmas'][1]['id']}/firmar", ops)

# PDF hecho en Python (ReportLab), sin Adobe ni navegador
r = c.get(f"/api/protocolos/{A}/pdf")
ok(r.status_code == 200 and r.content[:4] == b"%PDF", "PDF de revisión generado en Python")
lector = PdfReader(io.BytesIO(r.content))
ok(len(lector.pages) >= 2, f"PDF une formulario + imagen adjunta ({len(lector.pages)} páginas)")
texto = texto_pdf(r.content)
ok("ATF-RC-001" in texto and "Reservorio Cunahuiri" in texto and rc["planos"][0] in texto, "PDF muestra correlativo, entregable y planos")

# Paquete a KP + Hoja de Cargo
r = api("POST", "/paquetes", sup, json={"instancias": [A, B], "contacto_id": kp})
paq = api("GET", "/paquetes")[0]
ok(r["codigo"] == paq["codigo"] and paq["enlace_demo"], f"Paquete {r['codigo']} creado")
hc = c.get(f"/api/paquetes/{paq['id']}/hoja-cargo")
ok(hc.status_code == 200 and "ATF-RC-001" in texto_pdf(hc.content) and "ST.CA.PG.001.F007" in texto_pdf(hc.content)
   and "Reporte Topográfico" in texto_pdf(hc.content), "Hoja de Cargo con los registros del paquete (y el reporte topográfico)")
token = paq["enlace_demo"].split("t=")[1]
api("GET", "/revision/token-falso-que-no-existe-123", esperado=404)
ok(True, "Enlace falso: rechazado")
rev = api("GET", f"/revision/{token}")
ok(len(rev["items"]) == 2 and rev["contacto"]["empresa"] == "Knight Piésold" and "@" in rev["contacto"]["correo"],
   "Portal KP muestra los 2 protocolos (correo enmascarado)")
ok(c.get(rev["items"][0]["pdf_url"]).status_code == 200, "KP puede abrir el PDF desde el portal")

r = api("POST", f"/revision/{token}/decidir", esperado=400,
        json={"decisiones": [{"instancia_id": A, "decision": "aceptado"}], "codigo": "1", "firma": firma_png()})
ok("recepción" in r["error"], "Antes de revisar debe confirmar la recepción")
api("POST", f"/revision/{token}/codigo")
api("POST", f"/revision/{token}/recepcion", json={"codigo": ultimo_codigo(), "firma": firma_png()})
paq = next(p for p in api("GET", "/paquetes") if p["id"] == paq["id"])
ok(paq["firmado_recepcion"], "Confirmo recepción: KP firma la Hoja de Cargo al recibir")
api("POST", f"/revision/{token}/recepcion", esperado=400, json={"codigo": "1", "firma": firma_png()})
ok(True, "La recepción se confirma una sola vez")
decisiones = {"decisiones": [{"instancia_id": A, "decision": "aceptado"},
                             {"instancia_id": B, "decision": "rechazado", "motivo": "Falta ensayo"}], "firma": firma_png()}
r = api("POST", f"/revision/{token}/decidir", esperado=400, json={**decisiones, "codigo": "123456"})
ok("solicite el código" in r["error"], "Sin pedir código no se puede firmar")
api("POST", f"/revision/{token}/codigo")
codigo = ultimo_codigo()
r = api("POST", f"/revision/{token}/decidir", esperado=400, json={**decisiones, "codigo": "000000" if codigo != "000000" else "111111"})
ok("incorrecto" in r["error"], "Código incorrecto: rechazado")
r = api("POST", f"/revision/{token}/decidir", esperado=400, json={**decisiones, "codigo": codigo, "firma": ""})
ok("firma" in r["error"].lower(), "Sin firma dibujada: rechazado")
api("POST", f"/revision/{token}/decidir", esperado=400,
    json={**decisiones, "codigo": codigo, "decisiones": [{"instancia_id": B, "decision": "rechazado", "motivo": ""}]})
ok(True, "Rechazo sin motivo: no se acepta")
ok(rev["contacto"]["tiene_firma"] is False, "KP aún no tiene firma guardada")
api("POST", f"/revision/{token}/decidir", json={**decisiones, "codigo": codigo, "origen": "imagen", "guardar": True})
ok(api("GET", f"/protocolos/{A}")["inst"]["estado"] == "aprobado", "A aceptado y firmado por KP → Aprobado")
d = api("GET", f"/protocolos/{R}", cad)
ok(d["inst"]["estado"] == "aprobado" and d["firmas"][2]["estado"] == "firmado", "KP firmó el reporte topográfico junto con el protocolo")
ok(api("GET", f"/revision/{token}")["contacto"]["tiene_firma"], "KP cargó la imagen de su firma y la dejó guardada")
ok(api("GET", f"/protocolos/{B}")["inst"]["estado"] == "observado", "B rechazado → Observado (vuelve a Calidad)")
api("POST", f"/revision/{token}/decidir", esperado=400, json={**decisiones, "codigo": codigo})
ok(True, "El código sirve una sola vez")
api("POST", "/paquetes", sup, json={"instancias": [C], "contacto_id": kp})
t2 = api("GET", "/paquetes")[0]["enlace_demo"].split("t=")[1]
api("POST", f"/revision/{t2}/codigo")
api("POST", f"/revision/{t2}/recepcion", json={"codigo": ultimo_codigo(), "origen": "guardada"})
api("POST", f"/revision/{t2}/codigo")
api("POST", f"/revision/{t2}/decidir", json={"decisiones": [{"instancia_id": C, "decision": "aceptado"}],
                                            "codigo": ultimo_codigo(), "origen": "guardada"})
fc = next(f for f in api("GET", f"/protocolos/{C}")["firmas"] if f.get("contacto_id"))
ok(fc["estado"] == "firmado" and (fc["sello_url"] or "").startswith("data:image/png"), "En el siguiente paquete KP firma con su firma guardada")

# PDF final con constancia de firmas
r = c.get(f"/api/protocolos/{A}/pdf")
texto = texto_pdf(r.content)
ok(r.status_code == 200 and "revision" not in r.headers["content-disposition"], "PDF final firmado disponible")
ok("Constancia de firmas".upper() in texto.upper() and "revisor.kp@demo" in texto, "PDF final trae la constancia con el correo de KP")
ok(re.search(r"[0-9a-f]{64}", texto) is not None, "Constancia con código de verificación SHA-256")

# QC Index
r = c.get("/api/qc-index?entregable=RC")
wb = load_workbook(io.BytesIO(r.content))
ok("A. Fundac" in wb.sheetnames and "A. Relleno" in wb.sheetnames, "QC Index con las mismas hojas que el Excel de Calidad")
ws = wb["A. Fundac"]
fila = [cel.value for cel in ws[6]]
ok(fila[4] == "ST.CA.PC.001.F002 [10203002] - ATF-RC-001" and "Conforme" in fila, "QC Index: registro ATF-RC-001 Conforme")
ok(fila[10] == 19.51, f"QC Index calcula la longitud desde las progresivas ({fila[10]})")
ok(ws.max_row == 7, "QC Index filtrado por entregable (solo RC)")

# Control de cambios
api("PUT", f"/protocolos/{A}/valor", sup, 400, json={"cabecera": cab_a, "respuestas": resp_a, "conforme": True})
ok(True, "Aprobado = bloqueado")
api("POST", f"/protocolos/{A}/cambios", ops, 403, json={"motivo": "x", "contacto_id": kp})
ok(True, "Operaciones no solicita modificaciones")
api("POST", f"/protocolos/{A}/cambios", sup, json={"motivo": "Error en cota", "descripcion": "Corregir", "contacto_id": kp})
cambio = api("GET", f"/protocolos/{A}")["cambios"][0]
ok(cambio["etapa"] == "jefatura" and not cambio["enlace_demo"], "Paso 1: la solicitud va primero a la Jefatura de Calidad")
ok(not any("Solicitud de modificación" in x["asunto"] for x in api("GET", "/correos")), "KP todavía no recibe nada")
api("POST", f"/cambios/{cambio['id']}/decidir", sup, 403, json={"decision": "aceptado"})
ok(True, "Solo la Jefatura de Calidad valida la modificación")
api("POST", f"/cambios/{cambio['id']}/decidir", jefe, json={"decision": "aceptado"})
cambio = api("GET", f"/protocolos/{A}")["cambios"][0]
ok(cambio["etapa"] == "externo" and cambio["enlace_demo"], "Paso 2: validada por la Jefatura, pasa a KP")
ok(api("GET", f"/protocolos/{A}")["inst"]["estado"] == "aprobado", "Mientras KP no apruebe, sigue aprobado (V0)")
t_cambio = cambio["enlace_demo"].split("t=")[1]
api("POST", f"/revision/{t_cambio}/codigo")
api("POST", f"/revision/{t_cambio}/cambio", json={"decision": "aceptado", "codigo": ultimo_codigo()})
a = api("GET", f"/protocolos/{A}")
ok(a["inst"]["estado"] == "en_modificacion" and a["inst"]["version_actual"] == 1, "Paso 3: KP (con código) aprueba → se abre V1")
ok(a["inst"]["correlativo"] == "ATF-RC-001", "La V1 conserva el mismo correlativo")
ok(c.get(f"/api/protocolos/{A}/pdf?version=0").status_code == 200, "V0 y su PDF firmado se conservan")

# Permisos por rol
r = api("POST", "/protocolos", ops, 403, json={"tipo_id": fund["id"], "origen": "backlog"})
ok("rol" in r["error"], "Operaciones no puede crear protocolos (solo revisa y firma)")
api("GET", "/usuarios", sup, 403)
ok(True, "Solo un administrador ve la gestión de usuarios")
r = api("POST", "/usuarios", jefe, json={"nombre": "Topógrafo de prueba", "email": "topo@demo.local", "cargo": "Técnico",
                                          "rol": "topografia"})
ok(r["usuario"]["rol"] == "topografia" and r["usuario"]["debe_cambiar_clave"] == 1,
   "El administrador da de alta un usuario (debe cambiar su clave)")
topo = r["usuario"]["id"]
api("POST", "/usuarios", jefe, 400, json={"nombre": "Repetido", "email": "topo@demo.local", "rol": "topografia"})
ok(True, "No se repiten correos de usuario")
topografico = next(t_ for t_ in tipos if t_["codigo"] == "ST.CA.PC.001.F001")
api("POST", "/protocolos", topo, json={"tipo_id": topografico["id"], "origen": "solicitud_operaciones"})
r = api("POST", "/protocolos", topo, 403, json={"tipo_id": fund["id"], "origen": "solicitud_operaciones"})
ok("Topografía solo llena el Reporte Topográfico" in r["error"], "Topografía solo puede llenar el Reporte Topográfico")
api("PATCH", f"/usuarios/{topo}", jefe, json={"activo": False})
api("GET", "/tipos", topo, 401)
ok(True, "Usuario desactivado ya no entra")
api("PATCH", f"/usuarios/{jefe}", jefe, 400, json={"es_admin": False})
ok(True, "Un administrador no puede quitarse su propio acceso")
png = io.BytesIO()
Image.new("RGBA", (500, 200), (255, 255, 255, 0)).save(png, "PNG")
api("POST", f"/usuarios/{ops}/sello", sup, 403, files={"archivo": ("firma.png", png.getvalue(), "image/png")})
api("POST", f"/usuarios/{ops}/sello", ops, files={"archivo": ("firma.png", png.getvalue(), "image/png")})
ok(c.get(f"/api/usuarios/{ops}/sello").content[:4] == b"\x89PNG", "Cada firmante sube su firma escaneada (solo la suya)")
r = api("POST", "/contactos", jefe, json={"nombre": "Supervisora Hudbay", "email": "sup@hudbay.example", "cargo": "Supervisor",
                                           "empresa": "Hudbay"})
ok(r["empresa"] == "Hudbay", "El administrador registra firmantes externos")

# Configuración de Topografía (solo administradores)
from openpyxl import Workbook  # noqa: E402
wb = Workbook()
wb.active.append(["LOG DE REPORTES DE INSPECCIÓN DE LOS EQUIPOS DE MEDICIÓN Y/O ENSAYO"])
wb.active.append(["N° RIEME", "FECHA", "PROVEEDOR", "N° OC / OS", "N° GUÍA DE REMISIÓN", "EQUIPOS INSPECCIONADOS", "MARCA",
                  "MODELO", "N° SERIE / CÓDIGO", "ENTIDAD CALIBRADORA", "N° DE CERTIFICADO DE CALIBRACIÓN U OPERATIVIDAD",
                  "FECHA DE CALIBRACIÓN", "FRECUENCIA (meses)", "PRÓXIMA CALIBRACIÓN"])
wb.active.append(["RIEME-TOP-2027-01", None, "HUDBAY", "-", "-", "RECEPTOR GNSS", "TRIMBLE", "R12i", "NUEVO123", "GEO", "27-OG-0001",
                  "2027-03-01", 12, "2028-03-01"])
wb.active.append(["RIEME-TOP-2027-01", None, "HUDBAY", "-", "-", "NIVEL AUTOMATICO", "SPECTRA", "AL32A", "N1", "GEO", "x", None, 12, None])
xlsx = io.BytesIO()
wb.save(xlsx)
api("POST", "/topografia/equipos", cad, 403, files={"archivo": ("eme.xlsx", xlsx.getvalue())})
api("POST", "/topografia/equipos", jefe, files={"archivo": ("eme.xlsx", xlsx.getvalue())})
conf = api("GET", "/topografia/config", cad)
ok([e["modelo"] for e in conf["equipos"]] == ["R12i"], "El Excel de equipos del año reemplaza al anterior (sin nivel automático)")
ok(api("GET", f"/protocolos/{R}", cad)["inst"]["respuestas"]["tablas"]["equipos"][0]["modelo"] == "R980",
   "Los reportes ya hechos conservan su equipo")
api("POST", "/topografia/topografos", jefe, json={"nombre": "  Topógrafo   Nuevo "})
ok("Topógrafo Nuevo" in api("DELETE", "/topografia/topografos/Pamela Choque", jefe), "Se agregan y quitan topógrafos")
api("PUT", "/topografia/georreferencia", cad, 403, json={"punto": "X", "este": 1, "norte": 2, "elevacion": 3})
api("PUT", "/topografia/georreferencia", jefe, json={"punto": "BM-2", "este": "1", "norte": "2", "elevacion": "3"})
ok(api("GET", "/topografia/config", cad)["georreferencia"]["punto"] == "BM-2", "La georreferencia se puede modificar (otra unidad)")

print(f"\nTodo correcto: {pasos} verificaciones.")
